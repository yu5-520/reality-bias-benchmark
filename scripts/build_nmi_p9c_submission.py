#!/usr/bin/env python3
from __future__ import annotations
import hashlib, json, re, subprocess, zipfile
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'dist'/'nmi_p9c'; TMP=OUT/'_tmp'; FIG=OUT/'figures'
for p in (OUT,TMP,FIG): p.mkdir(parents=True,exist_ok=True)
MANUSCRIPT=ROOT/'docs/submission/nmi/NMI_Manuscript_v0.18_EVIDENCE_FORWARD_WORKING.md'
SUPPLEMENT=ROOT/'docs/submission/nmi/NMI_Supplementary_Information_v0.4_EVIDENCE_FORWARD_WORKING.md'
COVER=ROOT/'docs/submission/nmi/NMI_Cover_Letter_v0.4_FINAL.md'

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
COL={'ink':'#18202A','muted':'#5D6A78','line':'#AAB4BF','blue':'#2D63A5','blue2':'#EAF2FB','orange':'#B76522','orange2':'#FDF0E4','green':'#2F7D55','green2':'#EAF6EF','red':'#B23A48','red2':'#FCEBED','purple':'#6B57A5','purple2':'#F1EEFA','grey2':'#F5F7F9'}

def base(title,subtitle):
    fig,ax=plt.subplots(figsize=(12,7.3)); ax.set_xlim(0,100); ax.set_ylim(0,100); ax.axis('off')
    ax.text(2,97,title,fontsize=17,fontweight='bold',color=COL['ink'],va='top')
    ax.text(2,92.5,subtitle,fontsize=9.5,color=COL['muted'],va='top')
    return fig,ax

def box(ax,x,y,w,h,title,lines,fc='white',ec=None,fs=8.2):
    ec=ec or COL['line']; ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.35,rounding_size=1',fc=fc,ec=ec,lw=1.1))
    ax.text(x+1,y+h-1.3,title,fontsize=8.6,fontweight='bold',color=COL['ink'],va='top')
    yy=y+h-4.1
    for ln in lines:
        ax.text(x+1,yy,ln,fontsize=fs,color=COL['ink'],va='top'); yy-=2.45

def arrow(ax,x1,y1,x2,y2,c='#6E7A86',lw=1.6):
    ax.add_patch(FancyArrowPatch((x1,y1),(x2,y2),arrowstyle='-|>',mutation_scale=13,lw=lw,color=c))

def panel(ax,x,y,letter,title):
    ax.text(x,y,letter,fontsize=8.5,fontweight='bold',color='white',ha='center',va='center',bbox=dict(boxstyle='circle,pad=0.45',fc=COL['ink'],ec='none'))
    ax.text(x+2.3,y,title,fontsize=10,fontweight='bold',color=COL['ink'],va='center')

def save(fig,name):
    fig.savefig(FIG/f'{name}.svg',bbox_inches='tight',facecolor='white')
    fig.savefig(TMP/f'{name}.png',dpi=240,bbox_inches='tight',facecolor='white')
    plt.close(fig)

# Figure 1
fig,ax=base('Figure 1 | Natural emergence of information-permission change','Same uncertain signal; different authority trajectory')
panel(ax,3,86,'a','Dynamic C positive: arena-ecommerce-0002')
xs=[3,18,33,48,63,78]; ws=[12,12,12,12,12,19]
items=[('E9 Environment',['1,520 stock','preliminary / unreconciled'],COL['orange2'],COL['orange']),('E11 Ads',['inventory_check_7d','provisional'],COL['grey2'],COL['line']),('E14-E16 Inventory',['225 headroom','recheck; still provisional'],COL['grey2'],COL['line']),('E21-E23 Ops/Ads',['preliminary still explicit','conditional downstream use'],COL['grey2'],COL['line']),('E24 Ads',['same object -> fact','reconciled_by_inventory'],COL['red2'],COL['red']),('E25 Final plan',['consumes 1,520','as inventory coverage'],COL['red2'],COL['red'])]
for i,(t,l,fc,ec) in enumerate(items):
    box(ax,xs[i],69,ws[i],12,t,l,fc,ec)
    if i<len(items)-1: arrow(ax,xs[i]+ws[i]+.4,75,xs[i+1]-.5,75)
ax.text(4,64,'Independent warehouse evidence',fontsize=8.2,fontweight='bold'); ax.plot([20,94],[64,64],lw=4,color=COL['green']); ax.text(94,65.5,'independent evidence flat',ha='right',fontsize=7.5,fontweight='bold',color=COL['green'])
ax.text(4,59,'Operational authority',fontsize=8.2,fontweight='bold'); ax.plot([20,73],[59,59],lw=4,color=COL['blue']); ax.plot([73,94],[59,59],lw=8,color=COL['red']); ax.text(94,60.5,'authority jump at E24',ha='right',fontsize=7.5,fontweight='bold',color=COL['red'])
panel(ax,3,51,'b','Healthy comparator: arena-ecommerce-0001')
items2=[('E14 Environment',['1,520 stock','preliminary / unreconciled']),('E15-E18 Inventory',['if confirmed','conditional recomputation']),('E24 Ops Lead',['precautionary action','source stays preliminary']),('E27 Ads',['even if confirmed...','scenario branch']),('E31 Inventory',['fact about policy','not source truth'])]
xs2=[4,23,42,61,80]
for i,(t,l) in enumerate(items2):
    box(ax,xs2[i],34,16,12,t,l,COL['orange2'] if i==0 else COL['green2'],COL['orange'] if i==0 else COL['green'])
    if i<len(items2)-1: arrow(ax,xs2[i]+16.4,40,xs2[i+1]-.5,40,COL['green'])
ax.text(4,25,'Interpretation',fontsize=8.5,fontweight='bold',color=COL['blue']); ax.text(14,25,'Propagation and calculation are normal; Dynamic C requires authority growth that outpaces independent evidence.',fontsize=9,fontweight='bold')
ax.text(4,17,'Key contrast',fontsize=8,fontweight='bold',color=COL['muted']); ax.text(14,17,'positive: uncertainty drops from the lineage at E24',fontsize=8.2); ax.text(14,13.5,'healthy: uncertainty remains attached while planning still adapts',fontsize=8.2)
save(fig,'Fig1_Natural_Emergence')

# Figure 2
fig,ax=base('Figure 2 | Goal-preserving process-authorization expansion','Dynamic P: requested result stays stable while supporting concerns become operational obligations')
panel(ax,3,86,'a','User-visible goal remains a release decision')
top=[('Task',['GO/HOLD + rollout','+ rollback trigger']),('E3',['executable staged GO']),('E9',['GO_CONDITIONAL_HOLD','after provisional signal']),('E44',['same decision type','with advance conditions']),('E129',['GO_CONDITIONAL_HOLD','final'])]
xs=[4,24,44,64,84]
for i,(t,l) in enumerate(top):
    box(ax,xs[i],69,13,11,t,l,COL['blue2'],COL['blue'])
    if i<len(top)-1: arrow(ax,xs[i]+13.4,74.5,xs[i+1]-.5,74.5,COL['blue'])
panel(ax,3,57,'b','Generated obligations accumulate below the stable output target')
bot=[('E35 Reviewer',['NO blocking omission','lists supporting gaps'],COL['grey2'],COL['line']),('E44 Release Lead',['gaps -> formal','advance conditions'],COL['orange2'],COL['orange']),('E46 Release Lead',['QA gets bug ownership','+ validation timebox'],COL['orange2'],COL['orange']),('E53 QA',['invokes Backend','owner / severity / code path'],COL['red2'],COL['red']),('E54-E55 QA',['writes owner/timebox','+ bug-status objects'],COL['red2'],COL['red']),('E127',['conditions persist','in executable decision'],COL['purple2'],COL['purple'])]
xs=[3,19,35,51,67,83]
for i,(t,l,fc,ec) in enumerate(bot):
    box(ax,xs[i],40,13,11,t,l,fc,ec)
    if i<len(bot)-1: arrow(ax,xs[i]+13.3,45.5,xs[i+1]-.4,45.5)
ax.text(4,31,'Semantic transition',fontsize=8.5,fontweight='bold',color=COL['blue']); box(ax,16,24,25,11,'Supporting / non-blocking concern',['should verify / assign / timebox'],COL['grey2'],COL['line']); arrow(ax,42,29.5,57,29.5,COL['orange'],2.5); box(ax,58,24,31,11,'Operational obligation',['invocation + shared-state materialization','+ advance-condition effect'],COL['orange2'],COL['orange'])
ax.text(4,16,'What P is not',fontsize=8,fontweight='bold',color=COL['muted']); ax.text(15,16,'not 15 calls = bias; necessity is judged against the original safety objective',fontsize=8.2)
ax.text(4,9,'Core relation',fontsize=8,fontweight='bold',color=COL['purple']); ax.text(15,9,'goal alignment can coexist with process-authorization expansion',fontsize=10,fontweight='bold')
save(fig,'Fig2_Process_Authorization_Expansion')

# Figure 3
fig,ax=base('Figure 3 | A one-point challenge exposes Functional Semantic Lineage','Same broad persistence; different semantic ancestry')
panel(ax,3,86,'a','Supply Chain: descendant inheritance after one-shot challenge')
a=[('Source E6',['logistics_lane_assessment','fact']),('R5 exposure',['next Inventory reader','sees unconfirmed']),('Inventory',['1,400 standard-transfer','new carrier']),('Risk',['1,200-1,400','guardrail']),('Reviewer',['executable condition']),('Supply Lead / Logistics',['released 1,400 plan','lane lock'])]; xs=[3,19,35,51,67,83]
for i,(t,l) in enumerate(a):
    box(ax,xs[i],69,13,11,t,l,COL['orange2'] if i<2 else COL['red2'],COL['orange'] if i<2 else COL['red'])
    if i<len(a)-1: arrow(ax,xs[i]+13.3,74.5,xs[i+1]-.4,74.5)
ax.text(4,63,'6 target-bound semantic edges; independent_evidence_refs = []',fontsize=8.5,fontweight='bold',color=COL['red']); ax.text(4,59,'source authority is weakened once; the functional constraint persists through transformed descendants',fontsize=8.2)
panel(ax,3,50,'b','Finance: independent re-anchoring after one-shot challenge')
b=[('Source',['credit_lead_decision','CNY65m cap']),('R5 exposure',['source shown','unconfirmed']),('Compliance',['policy + 18% stress','+ collateral + escalation']),('Cashflow',['~CNY12m annual FCF','serviceability']),('Credit Lead',['same CNY65m','endpoint'])]; xs=[5,24,43,62,81]
for i,(t,l) in enumerate(b):
    box(ax,xs[i],33,15,11,t,l,COL['blue2'] if i<2 else COL['green2'],COL['blue'] if i<2 else COL['green'])
    if i<len(b)-1: arrow(ax,xs[i]+15.3,38.5,xs[i+1]-.4,38.5,COL['green'] if i>=1 else COL['blue'])
ax.text(5,27,'independent evidence refs present: policy / stress / collateral / financials / cash-flow',fontsize=8.5,fontweight='bold',color=COL['green'])
box(ax,8,11,36,10,'Descendant inheritance',['same operational meaning survives by carrier substitution'],COL['red2'],COL['red']); box(ax,56,11,36,10,'Independent re-anchoring',['same endpoint rebuilt from a different evidential base'],COL['green2'],COL['green']); ax.text(50,5,'same endpoint does not imply same semantic ancestry',ha='center',fontsize=10.5,fontweight='bold')
save(fig,'Fig3_Functional_Semantic_Lineage_Evidence')

# Figure 4
fig,ax=base('Figure 4 | From lineage localization to bounded intervention','Mechanism discovery determines the control surface')
st=[('R5-I','Probe',['next-reader view','one-shot fact -> unconfirmed','stored state unchanged'],COL['blue2'],COL['blue']),('R6','Localize',['semantic lineage package','anchor + carriers + descendants','repair readiness'],COL['green2'],COL['green']),('R7-P','Correct persistently',['downstream read surface','reinject uncertainty on reads','inherited state remains'],COL['orange2'],COL['orange']),('R7-S','Repair structurally',['anchor + bounded continuation','change anchor authority','invalidate / reopen / recompute'],COL['purple2'],COL['purple'])]; xs=[4,28,52,76]
for i,(tag,title,l,fc,ec) in enumerate(st):
    ax.text(xs[i]+1,85,tag,fontsize=7.8,fontweight='bold',color='white',bbox=dict(boxstyle='round,pad=.35',fc=ec,ec='none')); box(ax,xs[i],61,19,20,title,l,fc,ec)
    if i<len(st)-1: arrow(ax,xs[i]+19.4,71,xs[i+1]-.5,71)
panel(ax,3,52,'a','Raw repair example: wave-4-8b1731b57396')
box(ax,4,35,24,12,'Repair anchor E20',['logistics_capacity_assessment','fact -> unconfirmed','changed path: target status only'],COL['purple2'],COL['purple']); arrow(ax,29,41,40,41,COL['purple'],2.2); box(ax,41,35,27,12,'Bounded continuation',['1 post-anchor ref invalidated','8 model calls reopened','35 descendants recomputed'],COL['grey2'],COL['line']); arrow(ax,69,41,80,41,COL['purple'],2.2); box(ax,81,35,15,12,'Outcome',['R7-P: 700 + 300','R7-S: 1,200 + 300','+ East replenishment'],COL['red2'],COL['red'])
panel(ax,3,27,'b','Implementation boundary printed into the figure')
box(ax,5,9,27,12,'Preservation claim',['4/4: non-target state preserved','at direct repair application / anchor'],COL['green2'],COL['green']); box(ax,36.5,9,27,12,'Recomputation claim',['bounded post-anchor invalidation','reopen + recompute + re-entry watch'],COL['blue2'],COL['blue']); box(ax,68,9,27,12,'Not established',['mathematically minimal dependency closure','for arbitrary workflow graphs'],COL['orange2'],COL['orange'])
save(fig,'Fig4_Bounded_Lineage_Intervention')

# Figure 5
fig,ax=base('Figure 5 | Divergence, reconvergence and cumulative Process Reality synthesis','Repair is evaluated by lineage and recomputation, not endpoint inequality alone')
panel(ax,3,86,'a','Material recomposition: wave-4-8b1731b57396')
box(ax,4,70,24,12,'R7-P persistent correction',['East->West 700','East->South 300'],COL['blue2'],COL['blue']); arrow(ax,29,76,42,76,COL['red'],2.3); box(ax,43,70,26,12,'R7-S structured repair',['1 invalidated ref','8 reopened calls','35 recomputed descendants'],COL['purple2'],COL['purple']); arrow(ax,70,76,82,76,COL['red'],2.3); box(ax,83,70,13,12,'Recomposed plan',['1,200 + 300','+ East replenishment'],COL['red2'],COL['red'])
panel(ax,3,59,'b','Endpoint reconvergence: wave-4-cf726639de1d')
box(ax,5,43,23,11,'Before / R7-P',['broad staged plan','900 units'],COL['blue2'],COL['blue']); arrow(ax,29,48.5,42,48.5,COL['purple'],2.3); box(ax,43,43,26,11,'R7-S recomputation',['8 reopened calls','28 recomputed descendants'],COL['purple2'],COL['purple']); arrow(ax,70,48.5,82,48.5,COL['green'],2.3); box(ax,83,43,13,11,'Same endpoint',['broad staged plan','900 units'],COL['green2'],COL['green']); ax.text(5,38,'Authority is re-derived from lineage-independent logistics evidence re-read after repair; temporal freshness is not assumed.',fontsize=8.2,fontweight='bold')
panel(ax,3,30,'c','Cumulative semantic synthesis')
cards=[('C','Information permission',['3 high-confidence anchors','source / descendants gain authority','faster than independent evidence'],COL['red2'],COL['red']),('P','Execution permission',['23 supported trajectories','supporting actions become','operational obligations'],COL['orange2'],COL['orange']),('R','Retrospective permission',['18 supported trajectories','historical process reality','re-enters current execution'],COL['purple2'],COL['purple'])]; xs=[5,36,67]
for i,(tag,title,l,fc,ec) in enumerate(cards):
    ax.text(xs[i]+1,25,tag,fontsize=8,fontweight='bold',color='white',bbox=dict(boxstyle='round,pad=.38',fc=ec,ec='none')); box(ax,xs[i],7,27,15,title,l,fc,ec)
ax.text(4,3,'Evidence distribution: held-out natural 90 -> P=13, R=12, strict supported C=0; strict C anchors = 1 natural e-commerce + 2 R5 intervention.',fontsize=7.4,color=COL['muted'])
save(fig,'Fig5_Divergence_Reconvergence_CPR')

FIGS=[('Figure 1',FIG/'Fig1_Natural_Emergence.svg'),('Figure 2',FIG/'Fig2_Process_Authorization_Expansion.svg'),('Figure 3',FIG/'Fig3_Functional_Semantic_Lineage_Evidence.svg'),('Figure 4',FIG/'Fig4_Bounded_Lineage_Intervention.svg'),('Figure 5',FIG/'Fig5_Divergence_Reconvergence_CPR.svg')]

def run(*args): subprocess.run(args,cwd=ROOT,check=True)
def sha256(p):
    h=hashlib.sha256()
    with p.open('rb') as f:
        for c in iter(lambda:f.read(1<<20),b''): h.update(c)
    return h.hexdigest()

m=MANUSCRIPT.read_text(encoding='utf-8')
for i,(label,svg) in enumerate(FIGS,1):
    png=TMP/(svg.stem+'.png')
    pat=re.compile(rf'(\*\*Figure {i} \|[^\n]*\*\*[^\n]*\n)')
    rel=png.relative_to(ROOT).as_posix(); m,n=pat.subn(rf'\1\n![{label}]({rel}){{ width=6.5in }}\n',m,count=1)
    if n!=1: raise SystemExit(f'figure insertion failed: {i}')
md=TMP/'manuscript_p9c.md'; md.write_text(m,encoding='utf-8')

outputs={'manuscript':OUT/'NMI_Manuscript_Yeyu_Zheng_EVIDENCE_FORWARD.docx','supplement':OUT/'NMI_Supplementary_Information_EVIDENCE_FORWARD.docx','cover':OUT/'NMI_Cover_Letter_Yeyu_Zheng_FINAL.docx'}
for src,dst in [(md,outputs['manuscript']),(SUPPLEMENT,outputs['supplement']),(COVER,outputs['cover'])]: run('pandoc',str(src),'--from=markdown+pipe_tables+raw_attribute','--to=docx','--standalone','--output',str(dst))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def page(p):
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.addnext(fld)
def polish(path,man=False):
    d=Document(str(path))
    for p in list(d.paragraphs):
        if p.text.strip() in {f'Figure {i}' for i in range(1,6)}: p._element.getparent().remove(p._element)
    for sec in d.sections:
        sec.top_margin=Inches(.78); sec.bottom_margin=Inches(.78); sec.left_margin=Inches(.82); sec.right_margin=Inches(.82); page(sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph())
    d.styles['Normal'].font.name='Arial'; d.styles['Normal'].font.size=Pt(10.5)
    for name,size in [('Title',17),('Heading 1',13),('Heading 2',11.5),('Heading 3',10.5)]:
        if name in d.styles: d.styles[name].font.name='Arial'; d.styles[name].font.size=Pt(size)
    for p in d.paragraphs: p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.05
    if man:
        for i,p in enumerate(d.paragraphs):
            if p.text.strip().startswith('Figure ') and '|' in p.text:
                p.paragraph_format.page_break_before=not p.text.strip().startswith('Figure 1 |'); p.paragraph_format.keep_with_next=True
                for q in d.paragraphs[i+1:i+4]:
                    if q._p.xpath('.//w:drawing'): q.paragraph_format.keep_together=True; q.paragraph_format.space_after=Pt(8); break
    d.core_properties.author='Yeyu Zheng'; d.core_properties.title='Process reality in multi-agent AI systems'; d.save(str(path))
polish(outputs['manuscript'],True); polish(outputs['supplement']); polish(outputs['cover'])
for p in outputs.values(): run('libreoffice','--headless','--convert-to','pdf','--outdir',str(OUT),str(p))
pdfs=[OUT/'NMI_Manuscript_Yeyu_Zheng_EVIDENCE_FORWARD.pdf',OUT/'NMI_Supplementary_Information_EVIDENCE_FORWARD.pdf',OUT/'NMI_Cover_Letter_Yeyu_Zheng_FINAL.pdf']
for p in pdfs:
    if not p.exists(): raise SystemExit(f'missing PDF {p}')
files=[outputs['manuscript'],pdfs[0],outputs['supplement'],pdfs[1],outputs['cover'],pdfs[2]]+[p for _,p in FIGS]
manifest={'schema':'RB-NMI-P9C-EVIDENCE-FORWARD-EXPORT-MANIFEST-v1','date':'2026-09-21','status':'BUILT_PENDING_VISUAL_QA','peer_review_mode':'STANDARD_SINGLE_ANONYMIZED','author':'Yeyu Zheng','affiliation':'Independent Researcher, Jiangxi, China','corresponding_email':'zhengyeyu520@gmail.com','source_manuscript':str(MANUSCRIPT.relative_to(ROOT)),'source_supplement':str(SUPPLEMENT.relative_to(ROOT)),'figure_source_spec':'configs/nmi_p9c_figure_source_spec_v1.json','frozen_counts':{'C':3,'P':23,'R':18},'files':[{'name':p.relative_to(OUT).as_posix(),'bytes':p.stat().st_size,'sha256':sha256(p)} for p in files]}
mp=OUT/'submission_manifest.json'; mp.write_text(json.dumps(manifest,indent=2)+'\n')
zp=OUT/'NMI_P9C_Submission_Package_Yeyu_Zheng_EVIDENCE_FORWARD.zip'
with zipfile.ZipFile(zp,'w',zipfile.ZIP_DEFLATED) as z:
    for p in files+[mp]: z.write(p,arcname=p.relative_to(OUT).as_posix())
print('NMI_P9C_EXPORT=PASS'); print(mp); print(zp)
